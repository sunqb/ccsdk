"""RAG MVP regression tests."""
from __future__ import annotations

import asyncio
import nest_asyncio
from datetime import UTC, datetime, timedelta
from io import BytesIO

import httpx
import pytest
from fastapi.testclient import TestClient

nest_asyncio.apply()

from app.models.rag import RagRequestContext, RagSource, RagStreamRequest
from app.routers import rag as rag_router
from app.services.agent import AgentEvent
from app.services.rag import (
    LocalHashEmbeddingProvider,
    LocalVectorStore,
    OpenAICompatibleEmbeddingProvider,
    RagAgentRunner,
    RagAgentRunnerConfig,
    RagConcurrencyGuard,
    RagIngestionService,
    RagRetriever,
    SQLiteRagStateStore,
    RagToolService,
    TextChunker,
    TextDocumentParser,
    VectorStore,
    build_provider_info,
)
from app.services.rag.parser import (
    DocumentParser,
    HybridDocumentParser,
    LocalDocumentParser,
    MinerUDocumentParser,
)
from app.services.rag.mcp import (
    RAG_EMPTY_PROMPT_HINT,
    RAG_LOW_RELEVANCE_PROMPT_HINT,
    RAG_MIN_RELEVANCE_SCORE,
    build_rag_search_response,
)
from app.services.rag.tool_executor import RagToolExecutor
from app.services.rag.tool_schema import rag_direct_tools_schema
from app.services.rag.vector_store import SearchResult


def test_rag_stream_request_get_sources() -> None:
    request = RagStreamRequest(
        message="退款政策是什么？",
        knowledgeBaseId="kb_1",
        fileSetId="fs_1",
    )

    sources = request.get_sources()

    assert [(source.type, source.id) for source in sources] == [
        ("knowledge_base", "kb_1"),
        ("file_set", "fs_1"),
    ]


def test_text_parser_supports_txt_md_and_rejects_unknown_type() -> None:
    parser = TextDocumentParser()

    md = parser.parse_bytes(b"\xef\xbb\xbf# Title\n\nBody", filename="doc.md")
    txt = parser.parse_bytes("中文内容".encode(), filename="note.txt")

    assert md.mime_type == "text/markdown"
    assert md.text.startswith("# Title")
    assert txt.mime_type == "text/plain"
    assert txt.text == "中文内容"

    with pytest.raises(ValueError):
        parser.parse_bytes(b"unknown", filename="bad.xlsx")


def _build_docx_bytes(text: str) -> bytes:
    pytest.importorskip("docx")
    from docx import Document

    document = Document()
    document.add_paragraph(text)
    output = BytesIO()
    document.save(output)
    return output.getvalue()


def _build_pdf_bytes(text: str) -> bytes:
    canvas_module = pytest.importorskip("reportlab.pdfgen.canvas")

    output = BytesIO()
    canvas = canvas_module.Canvas(output)
    canvas.drawString(72, 720, text)
    canvas.save()
    return output.getvalue()


def test_parser_supports_docx_when_dependency_available() -> None:
    parser = TextDocumentParser()
    content = _build_docx_bytes("Refunds are available within 30 days.")

    document = parser.parse_bytes(content, filename="policy.docx")

    assert document.mime_type == (
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )
    assert "Refunds are available within 30 days." in document.text


def test_parser_supports_pdf_when_dependencies_available() -> None:
    pytest.importorskip("pypdf")
    parser = TextDocumentParser()
    content = _build_pdf_bytes("Refunds are available within 30 days.")

    document = parser.parse_bytes(content, filename="policy.pdf")

    assert document.mime_type == "application/pdf"
    assert "Refunds are available within 30 days." in document.text


def test_chunker_preserves_heading_metadata() -> None:
    parser = TextDocumentParser()
    document = parser.parse_bytes(
        b"# Root\n\nIntro\n\n## Child\n\nDetails",
        filename="guide.md",
    )
    chunks = TextChunker(chunk_size=100, chunk_overlap=10).chunk_document(document)

    assert chunks
    assert chunks[-1].metadata["filename"] == "guide.md"
    assert chunks[-1].metadata["heading_path"] == ["Root", "Child"]


@pytest.mark.asyncio
async def test_embedding_vector_store_search_and_file_set_scope() -> None:
    parser = TextDocumentParser()
    chunker = TextChunker(chunk_size=100, chunk_overlap=10)
    embedder = LocalHashEmbeddingProvider(dimensions=64)
    store = LocalVectorStore()

    refund_doc = parser.parse_bytes(
        b"Refunds are available within 30 days.",
        filename="refund.md",
        metadata={"file_set_id": "fs_refund"},
    )
    shipping_doc = parser.parse_bytes(
        b"Shipping usually takes 7 days.",
        filename="shipping.md",
        metadata={"file_set_id": "fs_shipping"},
    )
    chunks = [
        *chunker.chunk_document(refund_doc, source_file_id="file_refund"),
        *chunker.chunk_document(shipping_doc, source_file_id="file_shipping"),
    ]
    embeddings = await embedder.embed_documents([chunk.text for chunk in chunks])
    await store.upsert_chunks(chunks, embeddings)

    query = await embedder.embed_query("refund 30 days")
    results = await store.vector_search(
        query,
        sources=[{"type": "file_set", "id": "fs_refund"}],
        top_k=5,
    )

    assert results
    assert all(result.metadata["file_set_id"] == "fs_refund" for result in results)


@pytest.mark.asyncio
async def test_keyword_search_prioritizes_exact_terms_and_numbers() -> None:
    parser = TextDocumentParser()
    chunker = TextChunker(chunk_size=200, chunk_overlap=0)
    embedder = LocalHashEmbeddingProvider(dimensions=64)
    store = LocalVectorStore()

    exact_doc = parser.parse_bytes(
        b"Warranty plan ZX-900 covers battery replacement for exactly 18 months.",
        filename="warranty.md",
        metadata={"file_set_id": "fs_exact"},
    )
    generic_doc = parser.parse_bytes(
        b"General support covers account access and product onboarding.",
        filename="support.md",
        metadata={"file_set_id": "fs_exact"},
    )
    chunks = [
        *chunker.chunk_document(exact_doc, source_file_id="file_warranty"),
        *chunker.chunk_document(generic_doc, source_file_id="file_support"),
    ]
    embeddings = await embedder.embed_documents([chunk.text for chunk in chunks])
    await store.upsert_chunks(chunks, embeddings)

    results = await store.keyword_search(
        "ZX-900 18 months",
        sources=[{"type": "file_set", "id": "fs_exact"}],
        top_k=2,
    )

    assert results
    assert results[0].metadata["filename"] == "warranty.md"
    assert results[0].search_type == "keyword"
    assert "ZX-900" in results[0].text
    assert "18 months" in results[0].text


@pytest.mark.asyncio
async def test_openai_compatible_embedding_provider_posts_embeddings() -> None:
    captured: dict[str, object] = {}

    async def handler(request: httpx.Request) -> httpx.Response:
        captured["url"] = str(request.url)
        captured["authorization"] = request.headers.get("authorization")
        captured["payload"] = json_payload = request.read().decode()
        assert '"model":"embedding-model"' in json_payload
        assert '"input":["first","second"]' in json_payload
        return httpx.Response(
            200,
            json={
                "data": [
                    {"index": 1, "embedding": [0, 1.5]},
                    {"index": 0, "embedding": [2, 3]},
                ]
            },
        )

    client = httpx.AsyncClient(transport=httpx.MockTransport(handler))
    provider = OpenAICompatibleEmbeddingProvider(
        model="embedding-model",
        base_url="https://embedding.example/v1/embeddings/",
        api_key="test-key",
        client=client,
    )

    embeddings = await provider.embed_documents(["first", "second"])

    assert captured["url"] == "https://embedding.example/v1/embeddings"
    assert captured["authorization"] == "Bearer test-key"
    assert embeddings == [[2.0, 3.0], [0.0, 1.5]]
    await client.aclose()


@pytest.mark.asyncio
async def test_openai_compatible_embedding_provider_handles_errors() -> None:
    client = httpx.AsyncClient(
        transport=httpx.MockTransport(lambda request: httpx.Response(500, json={"error": "boom"}))
    )
    provider = OpenAICompatibleEmbeddingProvider(
        model="embedding-model",
        base_url="https://embedding.example/v1",
        client=client,
    )

    with pytest.raises(RuntimeError, match="HTTP 500"):
        await provider.embed_query("hello")

    await client.aclose()


@pytest.mark.asyncio
async def test_ingestion_ready_and_partial_ready() -> None:
    service = RagIngestionService(parser=TextDocumentParser())

    ready = await service.ingest_files(
        [("policy.md", b"# Refund\n\nRefunds are available within 30 days.")],
        conversation_id="conv_1",
    )

    assert ready.status == "ready"
    status = service.get_status(ready.file_set_id)
    assert status.progress == 100
    assert status.indexed_chunks == 1

    partial = await service.ingest_files([("bad.pdf", b"x"), ("ok.txt", b"hello")])
    assert partial.status == "partial_ready"
    assert [file.status for file in partial.files] == ["failed", "ready"]


@pytest.mark.asyncio
async def test_ingestion_persists_mysql_file_set_file_and_chunks() -> None:
    class FakeMySqlStore:
        def __init__(self) -> None:
            self.file_sets: list[dict[str, object]] = []
            self.file_set_updates: list[dict[str, object]] = []
            self.files: list[dict[str, object]] = []
            self.file_updates: list[dict[str, object]] = []
            self.chunk_batches: list[dict[str, object]] = []
            self.ingestion_jobs: list[dict[str, object]] = []
            self.ingestion_job_updates: list[dict[str, object]] = []

        async def save_file_set(self, **kwargs: object) -> None:
            self.file_sets.append(kwargs)

        async def update_file_set_status(self, **kwargs: object) -> None:
            self.file_set_updates.append(kwargs)

        async def save_file(self, **kwargs: object) -> None:
            self.files.append(kwargs)

        async def update_file_status(self, **kwargs: object) -> None:
            self.file_updates.append(kwargs)

        async def save_chunks(self, **kwargs: object) -> None:
            self.chunk_batches.append(kwargs)

        async def create_ingestion_job(self, **kwargs: object) -> None:
            self.ingestion_jobs.append(kwargs)

        async def update_ingestion_job(self, **kwargs: object) -> None:
            self.ingestion_job_updates.append(kwargs)

    mysql_store = FakeMySqlStore()
    service = RagIngestionService(parser=TextDocumentParser(), mysql_store=mysql_store)

    response = await service.ingest_files(
        [("policy.md", b"# Refund\n\nRefunds are available within 30 days.")],
        conversation_id="conv_mysql",
        metadata={"tenant_id": "tenant_1", "owner_id": "owner_1"},
    )

    assert response.status == "ready"
    assert mysql_store.file_sets[0]["file_set_id"] == response.file_set_id
    assert mysql_store.file_sets[0]["tenant_id"] == "tenant_1"
    assert mysql_store.files[0]["filename"] == "policy.md"
    assert mysql_store.file_set_updates[-1]["status"] == "ready"
    assert mysql_store.file_set_updates[-1]["indexed_chunks"] == 1
    assert mysql_store.chunk_batches
    assert len(mysql_store.chunk_batches[0]["chunks"]) == 1
    assert mysql_store.chunk_batches[0]["embedding_provider"] == "local"
    assert mysql_store.chunk_batches[0]["embedding_dimension"] == 256
    assert mysql_store.ingestion_jobs[0]["file_set_id"] == response.file_set_id
    assert mysql_store.ingestion_jobs[0]["status"] == "running"
    assert mysql_store.ingestion_job_updates[-1]["status"] == "succeeded"
    assert mysql_store.ingestion_job_updates[-1]["progress_percent"] == 100


@pytest.mark.asyncio
async def test_create_knowledge_base_updates_mysql_chunk_binding() -> None:
    class FakeMySqlStore:
        def __init__(self) -> None:
            self.file_sets: list[dict[str, object]] = []
            self.file_set_updates: list[dict[str, object]] = []
            self.files: list[dict[str, object]] = []
            self.file_updates: list[dict[str, object]] = []
            self.chunk_batches: list[dict[str, object]] = []
            self.knowledge_bases: list[dict[str, object]] = []
            self.kb_bindings: list[tuple[str, str]] = []
            self.chunk_bindings: list[dict[str, object]] = []

        async def save_file_set(self, **kwargs: object) -> None:
            self.file_sets.append(kwargs)

        async def update_file_set_status(self, **kwargs: object) -> None:
            self.file_set_updates.append(kwargs)

        async def save_file(self, **kwargs: object) -> None:
            self.files.append(kwargs)

        async def update_file_status(self, **kwargs: object) -> None:
            self.file_updates.append(kwargs)

        async def save_chunks(self, **kwargs: object) -> None:
            self.chunk_batches.append(kwargs)

        async def save_knowledge_base(self, **kwargs: object) -> None:
            self.knowledge_bases.append(kwargs)

        async def update_file_set_kb_binding(self, file_set_id: str, knowledge_base_id: str) -> None:
            self.kb_bindings.append((file_set_id, knowledge_base_id))

        async def update_chunks_knowledge_base(self, **kwargs: object) -> None:
            self.chunk_bindings.append(kwargs)

    mysql_store = FakeMySqlStore()
    service = RagIngestionService(parser=TextDocumentParser(), mysql_store=mysql_store)
    upload = await service.ingest_files([("policy.md", b"Refunds are available within 30 days.")])

    kb = await service.create_knowledge_base_from_file_set(
        file_set_id=upload.file_set_id,
        name="policy-kb",
    )

    assert mysql_store.knowledge_bases[0]["knowledge_base_id"] == kb.knowledge_base_id
    assert mysql_store.knowledge_bases[0]["embedding_provider"] == "local"
    assert mysql_store.kb_bindings == [(upload.file_set_id, kb.knowledge_base_id)]
    assert mysql_store.chunk_bindings[0]["file_set_id"] == upload.file_set_id
    assert mysql_store.chunk_bindings[0]["knowledge_base_id"] == kb.knowledge_base_id


@pytest.mark.asyncio
async def test_retriever_citations_and_read_chunk() -> None:
    service = RagIngestionService(parser=TextDocumentParser())
    response = await service.ingest_files(
        [("policy.md", b"# Refund\n\nRefunds are available within 30 days.")],
        conversation_id="conv_2",
    )
    retriever = RagRetriever(vector_store=service.get_vector_store(), embedder=service.embedder)
    sources = [RagSource(type="file_set", id=response.file_set_id)]

    results = await retriever.search("refund 30 days", sources=sources, top_k=2)
    chunks = await retriever.read_chunk(results[0].chunk_id, window=1)
    citations = retriever.build_citations(results)

    assert results
    assert chunks
    assert citations[0].source_name == "policy.md"
    assert citations[0].chunk_id == results[0].chunk_id


@pytest.mark.asyncio
async def test_rag_tool_service_enforces_scope_and_top_k() -> None:
    service = RagIngestionService(chunker=TextChunker(chunk_size=40, chunk_overlap=0))
    refund = await service.ingest_files(
        [("refund.md", b"Refund policy.\n\nRefunds are available within 30 days.")],
    )
    shipping = await service.ingest_files(
        [("shipping.md", b"Shipping policy.\n\nShipping usually takes 7 days.")],
    )
    retriever = RagRetriever(vector_store=service.get_vector_store(), embedder=service.embedder)
    tool_service = RagToolService(retriever=retriever)
    refund_context = RagRequestContext(
        requestId="req_scope",
        sources=[RagSource(type="file_set", id=refund.file_set_id)],
        topK=1,
    )
    shipping_results = await retriever.search(
        "shipping 7 days",
        sources=[RagSource(type="file_set", id=shipping.file_set_id)],
        top_k=1,
    )

    results = await tool_service.hybrid_search(
        query="policy days",
        context=refund_context,
        top_k=10,
    )
    blocked_chunks = await tool_service.read_chunk(
        chunk_id=shipping_results[0].chunk_id,
        context=refund_context,
        window=1,
    )

    assert len(results) == 1
    assert results[0].metadata["file_set_id"] == refund.file_set_id
    assert blocked_chunks == []


@pytest.mark.asyncio
async def test_hybrid_search_deduplicates_vector_and_keyword_matches() -> None:
    service = RagIngestionService(chunker=TextChunker(chunk_size=200, chunk_overlap=0))
    response = await service.ingest_files(
        [("policy.md", b"Refund code RMA-2026 allows refunds within exactly 30 days.")],
    )
    retriever = RagRetriever(vector_store=service.get_vector_store(), embedder=service.embedder)

    results = await retriever.search(
        "RMA-2026 30 days",
        sources=[RagSource(type="file_set", id=response.file_set_id)],
        top_k=5,
    )

    chunk_ids = [result.chunk_id for result in results]
    assert results
    assert len(chunk_ids) == len(set(chunk_ids))
    assert results[0].search_type in {"hybrid", "keyword", "vector"}


@pytest.mark.asyncio
async def test_hybrid_search_keeps_file_set_scope_for_keyword_matches() -> None:
    service = RagIngestionService(chunker=TextChunker(chunk_size=200, chunk_overlap=0))
    allowed = await service.ingest_files(
        [("allowed.md", b"Allowed policy mentions RMA-2026 but only for internal users.")],
    )
    blocked = await service.ingest_files(
        [("blocked.md", b"Blocked policy has stronger RMA-2026 refund details for 30 days.")],
    )
    retriever = RagRetriever(vector_store=service.get_vector_store(), embedder=service.embedder)
    context = RagRequestContext(
        requestId="req_keyword_scope",
        sources=[RagSource(type="file_set", id=allowed.file_set_id)],
        topK=5,
    )
    tool_service = RagToolService(retriever=retriever)

    results = await tool_service.hybrid_search(
        query="RMA-2026 30 days",
        context=context,
        top_k=5,
    )

    assert results
    assert all(result.metadata["file_set_id"] == allowed.file_set_id for result in results)
    assert all(result.metadata["file_set_id"] != blocked.file_set_id for result in results)


@pytest.mark.asyncio
async def test_rag_tool_service_get_file_outline_is_scoped() -> None:
    service = RagIngestionService(chunker=TextChunker(chunk_size=40, chunk_overlap=0))
    allowed = await service.ingest_files(
        [("allowed.md", b"# Alpha\n\nFirst section.\n\n# Beta\n\nSecond section.")],
    )
    blocked = await service.ingest_files([("blocked.md", b"# Secret\n\nHidden section.")])
    retriever = RagRetriever(vector_store=service.get_vector_store(), embedder=service.embedder)
    chunks = await service.get_vector_store().list_chunks(
        [RagSource(type="file_set", id=allowed.file_set_id)],
    )
    blocked_chunks = await service.get_vector_store().list_chunks(
        [RagSource(type="file_set", id=blocked.file_set_id)],
    )
    tool_service = RagToolService(retriever=retriever)
    context = RagRequestContext(
        requestId="req_outline",
        sources=[RagSource(type="file_set", id=allowed.file_set_id)],
        topK=5,
    )

    outline = await tool_service.get_file_outline(
        source_file_id=chunks[0].source_file_id or "",
        context=context,
    )
    blocked_outline = await tool_service.get_file_outline(
        source_file_id=blocked_chunks[0].source_file_id or "",
        context=context,
    )

    assert "mcp__rag__rag_get_file_outline" in rag_router.RAG_MCP_ALLOWED_TOOLS
    assert outline["found"] is True
    assert outline["sourceName"] == "allowed.md"
    assert outline["sections"]
    assert blocked_outline["found"] is False


def test_retriever_builds_multi_file_citation_metadata() -> None:
    results = [
        SearchResult(
            chunk_id="chunk_a",
            source_file_id="file_a",
            chunk_index=0,
            text="Alpha file says refunds are 30 days.",
            score=0.8,
            metadata={"filename": "alpha.md", "file_set_id": "fs_multi"},
            search_type="keyword",
        ),
        SearchResult(
            chunk_id="chunk_b",
            source_file_id="file_b",
            chunk_index=1,
            text="Beta file says refunds are 14 days.",
            score=0.7,
            metadata={"filename": "beta.md", "file_set_id": "fs_multi"},
            search_type="vector",
        ),
    ]

    citations = RagRetriever.build_citations(results)

    assert [citation.source_name for citation in citations] == ["alpha.md", "beta.md"]
    assert [citation.source_id for citation in citations] == ["file_a", "file_b"]
    assert citations[0].metadata["sourceFileId"] == "file_a"
    assert citations[1].metadata["chunkIndex"] == 1
    assert citations[0].metadata["searchType"] == "keyword"


def test_rag_mcp_search_response_empty_results_has_prompt_hint() -> None:
    response = build_rag_search_response("不存在的问题", [])

    assert response["results"] == []
    assert response["meta"]["status"] == "empty"
    assert response["meta"]["total"] == 0
    assert response["meta"]["aboveThreshold"] == 0
    assert response["meta"]["promptHint"] == RAG_EMPTY_PROMPT_HINT


def test_rag_mcp_search_response_low_relevance_warns_without_filtering() -> None:
    result = SearchResult(
        chunk_id="chunk_low",
        source_file_id="file_low",
        chunk_index=0,
        text="低相关片段",
        score=RAG_MIN_RELEVANCE_SCORE / 2,
        metadata={"filename": "low.md"},
    )

    response = build_rag_search_response("模糊问题", [result])

    assert response["meta"]["status"] == "low_relevance"
    assert response["meta"]["aboveThreshold"] == 0
    assert response["meta"]["promptHint"] == RAG_LOW_RELEVANCE_PROMPT_HINT
    assert "warning" in response["meta"]
    assert response["results"][0]["chunkId"] == "chunk_low"


def test_rag_mcp_search_response_ok_filters_below_threshold() -> None:
    low = SearchResult(chunk_id="chunk_low", text="low", score=0.1)
    high = SearchResult(chunk_id="chunk_high", text="high", score=0.9)

    response = build_rag_search_response("精确问题", [low, high], min_score=0.3)

    assert response["meta"]["status"] == "ok"
    assert response["meta"]["total"] == 2
    assert response["meta"]["aboveThreshold"] == 1
    assert [result["chunkId"] for result in response["results"]] == ["chunk_high"]


def test_local_vector_store_satisfies_vector_store_contract() -> None:
    store: VectorStore = LocalVectorStore()

    assert isinstance(store, LocalVectorStore)


def test_rag_direct_tool_schema_is_independent_from_router() -> None:
    tool_names = {tool["name"] for tool in rag_direct_tools_schema()}

    assert tool_names == {
        "rag_hybrid_search",
        "rag_read_chunk",
        "rag_list_sources",
        "rag_get_file_outline",
    }


def test_rag_agent_runner_config_keeps_direct_runtime_limits() -> None:
    runner = RagAgentRunner(
        config=RagAgentRunnerConfig(
            direct_timeout_seconds=30,
            direct_max_tokens=512,
            parse_only_max_tokens=128,
        )
    )

    assert runner.config.direct_timeout_seconds == 30
    assert runner.config.direct_max_tokens == 512
    assert runner.config.parse_only_max_tokens == 128


def test_parse_only_context_limit_can_be_disabled() -> None:
    from app.services.rag.agent_runner import _build_parsed_file_context

    rows = [{"filename": "long.txt", "parsed_text": "x" * 7000}]

    context_text, total_chars, truncated = _build_parsed_file_context(rows, max_tokens=None)

    assert truncated is False
    assert total_chars == len(context_text)
    assert "x" * 7000 in context_text


def test_parse_only_context_limit_truncates_when_configured() -> None:
    from app.services.rag.agent_runner import _build_parsed_file_context

    rows = [{"filename": "long.txt", "parsed_text": "x" * 7000}]

    context_text, total_chars, truncated = _build_parsed_file_context(rows, max_tokens=2)

    assert truncated is True
    assert total_chars == 6
    assert len(context_text) <= len("以下是用户上传的文件内容，请基于此回答问题：\n") + 6


@pytest.mark.asyncio
async def test_rag_tool_executor_uses_request_scoped_rag_tools() -> None:
    service = RagIngestionService(chunker=TextChunker(chunk_size=200, chunk_overlap=0))
    response = await service.ingest_files(
        [("policy.md", b"Executor says refund code EXEC-2026 allows 30 day refunds.")],
    )
    retriever = RagRetriever(vector_store=service.get_vector_store(), embedder=service.embedder)
    executor = RagToolExecutor(tool_service=RagToolService(retriever=retriever))
    context = RagRequestContext(
        requestId="req_executor",
        sources=[RagSource(type="file_set", id=response.file_set_id)],
        topK=3,
    )

    results = await executor.execute(
        name="rag_hybrid_search",
        tool_input={"query": "EXEC-2026 30 day refunds", "top_k": 10},
        context=context,
    )

    assert results
    assert len(results) <= 3
    assert results[0]["metadata"]["file_set_id"] == response.file_set_id


@pytest.mark.asyncio
async def test_rag_state_store_restores_knowledge_base_after_restart(tmp_path) -> None:
    state_store = SQLiteRagStateStore(tmp_path / "rag.db")
    service = RagIngestionService(
        chunker=TextChunker(chunk_size=200, chunk_overlap=0),
        state_store=state_store,
    )
    response = await service.ingest_files(
        [("policy.md", b"Persistent KB says refund code PERSIST-2026 allows 30 day refunds.")],
        conversation_id="conv_persist",
    )
    kb = await service.create_knowledge_base_from_file_set(
        file_set_id=response.file_set_id,
        name="Persistent Policy KB",
        tenant_id="tenant_persist",
    )

    restarted = RagIngestionService(
        chunker=TextChunker(chunk_size=200, chunk_overlap=0),
        state_store=SQLiteRagStateStore(tmp_path / "rag.db"),
    )
    retriever = RagRetriever(vector_store=restarted.get_vector_store(), embedder=restarted.embedder)
    results = await retriever.search(
        "PERSIST-2026 30 day refunds",
        sources=[
            RagSource(
                type="knowledge_base",
                id=kb.knowledge_base_id,
                metadata={"tenantId": "tenant_persist"},
            )
        ],
        top_k=3,
    )

    assert restarted.get_status(response.file_set_id).status == "ready"
    assert restarted.list_knowledge_bases(tenant_id="tenant_persist")[0].knowledge_base_id == kb.knowledge_base_id
    assert restarted.get_stats()["stateStore"] == "sqlite"
    assert results
    assert results[0].metadata["knowledge_base_id"] == kb.knowledge_base_id


def test_rag_files_upload_and_status_endpoint(monkeypatch: pytest.MonkeyPatch) -> None:
    from app import auth
    from app.main import app

    service = RagIngestionService(parser=TextDocumentParser())
    monkeypatch.setattr(auth, "get_api_key", lambda: None)
    monkeypatch.setattr(rag_router, "rag_ingestion_service", service)

    client = TestClient(app)
    response = client.post(
        "/rag/files",
        files=[("files", ("policy.md", b"# Policy\n\nRefunds within 30 days.", "text/markdown"))],
        data={"conversationId": "conv_http", "metadata": '{"owner":"test"}'},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "ready"
    assert body["conversationId"] == "conv_http"

    status = client.get(f"/rag/files/{body['fileSetId']}/status")
    assert status.status_code == 200
    assert status.json()["progress"] == 100

    missing = client.get("/rag/files/missing/status")
    assert missing.status_code == 404

    bad_metadata = client.post(
        "/rag/files",
        files=[("files", ("policy.md", b"x", "text/markdown"))],
        data={"metadata": "[]"},
    )
    assert bad_metadata.status_code == 400

    invalid_metadata = client.post(
        "/rag/files",
        files=[("files", ("policy.md", b"x", "text/markdown"))],
        data={"metadata": "{"},
    )
    assert invalid_metadata.status_code == 400


def test_rag_files_upload_accepts_repeated_file_field(monkeypatch: pytest.MonkeyPatch) -> None:
    from app import auth
    from app.main import app

    service = RagIngestionService(parser=TextDocumentParser())
    monkeypatch.setattr(auth, "get_api_key", lambda: None)
    monkeypatch.setattr(rag_router, "rag_ingestion_service", service)

    client = TestClient(app)
    response = client.post(
        "/rag/files",
        files=[
            ("file", ("refund.txt", b"Refunds within 30 days.", "text/plain")),
            ("file", ("shipping.txt", b"Shipping within 7 days.", "text/plain")),
        ],
        data={"conversationId": "conv_repeated_file"},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "ready"
    assert [file["filename"] for file in body["files"]] == ["refund.txt", "shipping.txt"]


def test_rag_files_upload_accepts_docx_and_pdf_when_dependencies_available(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    pytest.importorskip("docx")
    pytest.importorskip("pypdf")
    pytest.importorskip("reportlab.pdfgen.canvas")
    from app import auth
    from app.main import app

    service = RagIngestionService(parser=TextDocumentParser())
    monkeypatch.setattr(auth, "get_api_key", lambda: None)
    monkeypatch.setattr(rag_router, "rag_ingestion_service", service)

    client = TestClient(app)
    response = client.post(
        "/rag/files",
        files=[
            (
                "files",
                (
                    "policy.docx",
                    _build_docx_bytes("DOCX refunds within 30 days."),
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                ),
            ),
            ("files", ("policy.pdf", _build_pdf_bytes("PDF refunds within 30 days."), "application/pdf")),
        ],
        data={"conversationId": "conv_http_docs"},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["status"] == "ready"
    assert [file["mimeType"] for file in body["files"]] == [
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "application/pdf",
    ]


def test_rag_query_missing_sources_returns_400(monkeypatch: pytest.MonkeyPatch) -> None:
    from app import auth
    from app.main import app

    monkeypatch.setattr(auth, "get_api_key", lambda: None)
    client = TestClient(app)

    response = client.post("/rag/query", json={"message": "no sources"})

    assert response.status_code == 400
    assert response.json()["detail"]["code"] == "missing_sources"


# ------------------------------------------------------------------
# Parser abstraction tests
# ------------------------------------------------------------------

def test_local_document_parser_metadata_contains_parser_field() -> None:
    parser = LocalDocumentParser()
    doc = parser.parse_bytes(b"# Hello\n\nWorld", filename="test.md")

    assert doc.metadata["parser"] == "local"
    assert doc.metadata["extension"] == ".md"
    assert isinstance(parser.supported_extensions, set)
    assert ".md" in parser.supported_extensions
    assert ".txt" in parser.supported_extensions
    assert ".pdf" in parser.supported_extensions
    assert ".docx" in parser.supported_extensions


def test_local_document_parser_rejects_unsupported_types() -> None:
    parser = LocalDocumentParser()

    with pytest.raises(ValueError, match="Unsupported local document type"):
        parser.parse_bytes(b"data", filename="file.xlsx")


@pytest.mark.asyncio
async def test_mineru_parser_returns_markdown_on_success() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        assert request.url.path == "/parse"
        return httpx.Response(
            200,
            json={
                "markdown": "# 第一章 总则\n\n合同有效期为三年。",
                "page_count": 12,
                "document_id": "doc_abc123",
                "version": "1.2.3",
            },
        )

    client = httpx.Client(transport=httpx.MockTransport(handler))
    parser = MinerUDocumentParser(
        base_url="https://mineru.internal.example",
        api_key="secret-key",
        timeout_seconds=60.0,
        client=client,
    )

    doc = parser.parse_bytes(
        b"fake pdf bytes",
        filename="contract.pdf",
    )

    assert doc.text.startswith("# 第一章")
    assert "合同有效期为三年。" in doc.text
    assert doc.metadata["parser"] == "mineru"
    assert doc.metadata["pageCount"] == 12
    assert doc.metadata["documentId"] == "doc_abc123"
    assert doc.metadata["parserVersion"] == "1.2.3"
    assert doc.metadata["extension"] == ".pdf"
    client.close()


@pytest.mark.asyncio
async def test_mineru_parser_handles_wrapped_response() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(
            200,
            json={
                "content": "# 合同条款\n\n付款方式为银行转账。",
                "pageCount": 5,
            },
        )

    client = httpx.Client(transport=httpx.MockTransport(handler))
    parser = MinerUDocumentParser(
        base_url="https://mineru.internal.example",
        client=client,
    )

    doc = parser.parse_bytes(b"fake docx bytes", filename="agreement.docx")

    assert "合同条款" in doc.text
    assert doc.metadata["pageCount"] == 5
    client.close()


@pytest.mark.asyncio
async def test_mineru_parser_raises_on_http_error() -> None:
    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(500, json={"error": "internal error"})

    client = httpx.Client(transport=httpx.MockTransport(handler))
    parser = MinerUDocumentParser(
        base_url="https://mineru.internal.example",
        client=client,
    )

    with pytest.raises(ValueError, match="MinerU returned HTTP 500"):
        parser.parse_bytes(b"fake", filename="doc.pdf")

    client.close()


@pytest.mark.asyncio
async def test_mineru_parser_falls_back_to_local_on_failure() -> None:
    pytest.importorskip("reportlab.pdfgen.canvas")
    def handler(request: httpx.Request) -> httpx.Response:
        return httpx.Response(500, json={"error": "service unavailable"})

    client = httpx.Client(transport=httpx.MockTransport(handler))
    parser = MinerUDocumentParser(
        base_url="https://mineru.internal.example",
        fallback_to_local=True,
        client=client,
    )

    real_pdf_bytes = _build_pdf_bytes("Fallback PDF content via local parser.")

    doc = parser.parse_bytes(real_pdf_bytes, filename="fallback.pdf")

    assert doc.metadata["parser"] == "local"
    assert doc.metadata["parserFallbackFrom"] == "mineru"
    assert "Fallback PDF content" in doc.text
    client.close()


def test_hybrid_document_parser_routes_to_local_for_txt_md() -> None:
    parser = HybridDocumentParser()

    assert ".txt" in parser.supported_extensions
    assert ".md" in parser.supported_extensions

    doc = parser.parse_bytes(b"Hello world", filename="note.txt")
    assert doc.metadata["parser"] == "local"


def test_hybrid_document_parser_rejects_pdf_without_mineru() -> None:
    parser = HybridDocumentParser()

    with pytest.raises(ValueError, match="MinerU is not configured"):
        parser.parse_bytes(b"fake", filename="doc.pdf")


def test_hybrid_document_parser_with_mineru_config() -> None:
    parser = HybridDocumentParser(
        mineru_base_url="https://mineru.internal.example",
        mineru_api_key="key",
        mineru_timeout_seconds=30.0,
        mineru_fallback_to_local=False,
    )

    assert ".pdf" in parser.supported_extensions
    assert ".docx" in parser.supported_extensions
    assert ".txt" in parser.supported_extensions


def test_document_parser_protocol_is_runtime_checkable() -> None:
    local = LocalDocumentParser()
    hybrid = HybridDocumentParser()

    assert isinstance(local, DocumentParser)
    assert isinstance(hybrid, DocumentParser)


def test_ingestion_service_builds_parser_from_config_local(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr("app.config.settings.file_parser_provider", "local")
    service = RagIngestionService()

    assert isinstance(service.parser, LocalDocumentParser)
    assert service.parser.supported_extensions == {".txt", ".md", ".pdf", ".docx"}



def test_rag_query_returns_answer_and_citations_with_fake_agent(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from app import auth
    from app.main import app

    service = RagIngestionService()

    async def ingest() -> str:
        response = await service.ingest_files(
            [("policy.md", b"# Refund\n\nRefunds are available within 30 days.")],
        )
        return response.file_set_id

    file_set_id = asyncio.run(ingest())
    retriever = RagRetriever(vector_store=service.get_vector_store(), embedder=service.embedder)
    monkeypatch.setattr(auth, "get_api_key", lambda: None)
    monkeypatch.setattr(rag_router, "rag_tool_service", RagToolService(retriever=retriever))
    monkeypatch.setattr(rag_router, "create_rag_mcp_server", lambda *args, **kwargs: object())
    captured: dict[str, object] = {}

    async def fake_query_stream(**kwargs):  # noqa: ANN202
        captured.update(kwargs)
        yield AgentEvent(
            type="content_block_delta",
            subtype="text_delta",
            data={"text": "退款期限是 30 天。"},
        )

    monkeypatch.setattr(rag_router.agent_service, "query_stream", fake_query_stream)

    client = TestClient(app)
    response = client.post(
        "/rag/query",
        json={"message": "退款期限是多久？", "fileSetId": file_set_id},
    )

    assert response.status_code == 200
    body = response.json()
    assert body["answer"] == "退款期限是 30 天。"
    assert body["citations"]
    assert body["citations"][0]["sourceName"] == "policy.md"
    assert body["usage"]["retrieval"]["matchedChunks"] >= 1
    assert body["usage"]["retrieval"]["provider"] == "local"
    assert body["usage"]["agent"]["outputChars"] == len("退款期限是 30 天。")
    assert "policy.md" in captured["prompt"]


def test_rag_agent_stream_exposes_rag_tools_to_native_agent(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from app import auth
    from app.main import app

    service = RagIngestionService()

    async def ingest() -> str:
        response = await service.ingest_files(
            [("skills.md", b"# Task\n\nThis document should be translated and polished.")],
            conversation_id="conv_agent_tool",
        )
        return response.file_set_id

    file_set_id = asyncio.run(ingest())
    retriever = RagRetriever(vector_store=service.get_vector_store(), embedder=service.embedder)
    monkeypatch.setattr(auth, "get_api_key", lambda: None)
    monkeypatch.setattr(rag_router, "rag_tool_service", RagToolService(retriever=retriever))
    captured: dict[str, object] = {}

    async def fake_query_stream(**kwargs):  # noqa: ANN202
        captured.update(kwargs)
        yield AgentEvent(
            type="content_block_delta",
            subtype="text_delta",
            data={"text": "可以由 Claude Agent 原生 Skills 判断，文件更适合 translate 或 text-optimize。"},
        )

    monkeypatch.setattr(rag_router.agent_service, "query_stream", fake_query_stream)

    client = TestClient(app)
    response = client.post(
        "/rag/agent/stream",
        json={
            "message": "目前有哪些技能？然后这个文件里面应该用哪个技能合适",
            "conversationId": "conv_agent_tool",
            "fileSetId": file_set_id,
        },
    )

    rendered = response.text
    assert response.status_code == 200
    assert "event: result" in rendered
    assert "claude_sdk" in rendered or "agent_tool" in rendered
    assert captured["prompt"] == "目前有哪些技能？然后这个文件里面应该用哪个技能合适"
    assert captured.get("allowed_tools") == []
    assert captured["mcp_servers"]
    assert "rag" in captured["mcp_servers"]
    assert captured["cwd"] == rag_router.settings.work_dir


def test_knowledge_base_endpoints_create_list_delete(monkeypatch: pytest.MonkeyPatch) -> None:
    from app import auth
    from app.main import app

    service = RagIngestionService()

    async def ingest() -> str:
        response = await service.ingest_files(
            [("policy.md", b"# Refund\n\nRefunds are available within 30 days.")],
        )
        return response.file_set_id

    file_set_id = asyncio.run(ingest())
    monkeypatch.setattr(auth, "get_api_key", lambda: None)
    monkeypatch.setattr(rag_router, "rag_ingestion_service", service)

    client = TestClient(app)
    create_response = client.post(
        "/rag/knowledge-bases",
        json={
            "name": "Policy KB",
            "description": "Refund policy",
            "sourceFileSetId": file_set_id,
            "tenantId": "tenant_1",
            "ownerId": "owner_1",
            "apiKeyId": "key_1",
            "metadata": {"domain": "support"},
        },
    )

    assert create_response.status_code == 200
    kb = create_response.json()
    assert kb["knowledgeBaseId"].startswith("kb_")
    assert kb["sourceFileSetId"] == file_set_id
    assert kb["tenantId"] == "tenant_1"

    list_response = client.get("/rag/knowledge-bases", params={"tenantId": "tenant_1"})
    assert list_response.status_code == 200
    assert [item["knowledgeBaseId"] for item in list_response.json()["knowledgeBases"]] == [
        kb["knowledgeBaseId"]
    ]

    delete_response = client.delete(f"/rag/knowledge-bases/{kb['knowledgeBaseId']}")
    assert delete_response.status_code == 200
    assert delete_response.json()["status"] == "deleted"

    missing_delete = client.delete(f"/rag/knowledge-bases/{kb['knowledgeBaseId']}")
    assert missing_delete.status_code == 404


@pytest.mark.asyncio
async def test_knowledge_base_id_retrieval_and_permission_scope() -> None:
    service = RagIngestionService(chunker=TextChunker(chunk_size=200, chunk_overlap=0))
    response = await service.ingest_files(
        [("policy.md", b"Policy KB says refund code KB-2026 allows 30 day refunds.")],
    )
    kb = await service.create_knowledge_base_from_file_set(
        file_set_id=response.file_set_id,
        name="Policy KB",
        tenant_id="tenant_allowed",
        owner_id="owner_allowed",
        api_key_id="key_allowed",
    )
    retriever = RagRetriever(vector_store=service.get_vector_store(), embedder=service.embedder)

    allowed_results = await retriever.search(
        "KB-2026 30 day refunds",
        sources=[
            RagSource(
                type="knowledge_base",
                id=kb.knowledge_base_id,
                metadata={"tenantId": "tenant_allowed", "ownerId": "owner_allowed"},
            )
        ],
        top_k=3,
    )
    blocked_results = await retriever.search(
        "KB-2026 30 day refunds",
        sources=[
            RagSource(
                type="knowledge_base",
                id=kb.knowledge_base_id,
                metadata={"tenantId": "tenant_blocked"},
            )
        ],
        top_k=3,
    )

    assert allowed_results
    assert all(result.metadata["knowledge_base_id"] == kb.knowledge_base_id for result in allowed_results)
    assert all(result.metadata["tenant_id"] == "tenant_allowed" for result in allowed_results)
    assert blocked_results == []


@pytest.mark.asyncio
async def test_cleanup_expired_file_sets_skips_persistent_knowledge_base_chunks() -> None:
    service = RagIngestionService(chunker=TextChunker(chunk_size=200, chunk_overlap=0))
    expired = await service.ingest_files([("expired.md", b"Temporary expired content.")])
    persisted = await service.ingest_files([("persisted.md", b"Persistent KB content.")])
    kb = await service.create_knowledge_base_from_file_set(
        file_set_id=persisted.file_set_id,
        name="Persistent KB",
    )
    service._records[expired.file_set_id].expires_at = datetime.now(UTC) - timedelta(hours=1)
    service._records[persisted.file_set_id].expires_at = datetime.now(UTC) - timedelta(hours=1)

    cleaned = await service.cleanup_expired_file_sets()
    expired_chunks = await service.get_vector_store().list_chunks(
        [RagSource(type="file_set", id=expired.file_set_id)]
    )
    persisted_chunks = await service.get_vector_store().list_chunks(
        [RagSource(type="knowledge_base", id=kb.knowledge_base_id)]
    )

    assert cleaned == 1
    assert expired.file_set_id not in service._records
    assert expired_chunks == []
    assert persisted_chunks


def test_local_reranker_prioritizes_lexical_evidence() -> None:
    noisy = SearchResult(chunk_id="chunk_noisy", text="Generic account onboarding.", score=0.55)
    relevant = SearchResult(
        chunk_id="chunk_relevant",
        text="Refund code RMA-2026 allows refunds within 30 days.",
        score=0.05,
    )

    results = RagRetriever.rerank_results("RMA-2026 refund", [noisy, relevant])

    assert results[0].chunk_id == "chunk_relevant"
    assert results[0].search_type == "rerank"
    assert "rerankScore" in results[0].metadata


def test_query_rewrite_adds_domain_expansions() -> None:
    variants = RagRetriever.rewrite_query("return policy")

    assert variants[0] == "return policy"
    assert any("refund" in variant and "reimburse" in variant for variant in variants[1:])


@pytest.mark.asyncio
async def test_context_expansion_includes_neighboring_chunks() -> None:
    service = RagIngestionService(chunker=TextChunker(chunk_size=35, chunk_overlap=0))
    response = await service.ingest_files(
        [("manual.md", b"Intro paragraph.\n\nRefund terms are here.\n\nFinal notes.")],
    )
    chunks = await service.get_vector_store().list_chunks(
        [RagSource(type="file_set", id=response.file_set_id)]
    )
    retriever = RagRetriever(vector_store=service.get_vector_store(), embedder=service.embedder)
    anchor = chunks[1]

    expanded = await retriever.expand_results_with_context(
        [
            SearchResult(
                chunk_id=anchor.chunk_id,
                source_file_id=anchor.source_file_id,
                chunk_index=anchor.chunk_index,
                text=anchor.text,
                score=1.0,
                metadata=anchor.metadata,
                chunk=anchor,
            )
        ],
        sources=[RagSource(type="file_set", id=response.file_set_id)],
        window=1,
        top_k=3,
    )

    assert len(expanded) == 3
    assert {result.search_type for result in expanded} >= {"vector", "context"}
    assert any(result.metadata.get("contextAnchorChunkId") == anchor.chunk_id for result in expanded)


def test_rag_query_low_confidence_returns_insufficient_without_agent(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    from app import auth
    from app.main import app

    service = RagIngestionService()

    async def ingest() -> str:
        response = await service.ingest_files([("policy.md", b"Shipping takes seven days.")])
        return response.file_set_id

    file_set_id = asyncio.run(ingest())
    retriever = RagRetriever(vector_store=service.get_vector_store(), embedder=service.embedder)
    monkeypatch.setattr(auth, "get_api_key", lambda: None)
    monkeypatch.setattr(rag_router, "rag_tool_service", RagToolService(retriever=retriever))

    async def unexpected_query_stream(**kwargs):  # noqa: ANN202, ARG001
        raise AssertionError("low-confidence /rag/query must not call the agent")

    monkeypatch.setattr(rag_router.agent_service, "query_stream", unexpected_query_stream)

    client = TestClient(app)
    response = client.post(
        "/rag/query",
        json={
            "message": "ZX-999 unicorn refund clause?",
            "fileSetId": file_set_id,
            "options": {"minConfidence": 1.0, "lowConfidenceStrategy": "insufficient_context"},
        },
    )

    assert response.status_code == 200
    body = response.json()
    assert "不足" in body["answer"] or "没有找到" in body["answer"]
    assert body["usage"]["retrieval"]["confidence"] < 1.0
    assert body["usage"]["verification"]["reasons"]


def test_multi_query_builds_multiple_variants() -> None:
    variants = RagRetriever.build_query_variants(
        "return policy refund",
        query_rewrite=True,
        multi_query=True,
    )

    assert variants[0] == "return policy refund"
    assert len(variants) >= 3


@pytest.mark.asyncio
async def test_hybrid_search_records_retrieval_trace() -> None:
    service = RagIngestionService(chunker=TextChunker(chunk_size=80, chunk_overlap=0))
    response = await service.ingest_files(
        [("policy.md", b"Refund within 30 days.\n\nPayment invoice terms apply.")],
    )
    from app.services.rag.retriever import RetrievalTrace

    retriever = RagRetriever(vector_store=service.get_vector_store(), embedder=service.embedder)
    trace = RetrievalTrace(query="refund 30 days", variants=[], retrieve_top_k=20, final_top_k=3)
    results = await retriever.hybrid_search(
        "refund 30 days",
        sources=[RagSource(type="file_set", id=response.file_set_id)],
        retrieve_top_k=20,
        final_top_k=3,
        query_rewrite=True,
        multi_query=True,
        rerank=True,
        trace=trace,
    )

    assert results
    assert trace.stages
    assert any(stage["stage"] == "merge" for stage in trace.stages)


def test_chunk_metadata_includes_parent_chunk_fields() -> None:
    parser = TextDocumentParser()
    chunker = TextChunker(chunk_size=120, chunk_overlap=20)
    document = parser.parse_bytes(
        b"# Refund Policy\n\nRefunds are allowed within 30 days.",
        filename="policy.md",
    )
    chunks = chunker.chunk_document(document, source_file_id="file_1")

    assert chunks
    metadata = chunks[0].metadata
    assert metadata.get("parentChunkId")
    assert metadata.get("documentId")
    assert metadata.get("headingPath") == ["Refund Policy"]


def test_answer_verifier_detects_unsupported_answer() -> None:
    from app.services.rag.answer_verifier import rag_answer_verifier

    results = [
        SearchResult(
            chunk_id="chunk_1",
            text="Shipping takes seven days.",
            score=0.9,
        )
    ]
    verification = rag_answer_verifier.verify_answer(
        query="refund within 30 days",
        answer="Refunds are allowed within 30 days.",
        citations=[],
        results=results,
        min_alignment=0.6,
    )

    assert verification.status != "ok"
    assert "answer_not_supported_by_evidence" in verification.reasons


def test_retrieval_evaluation_harness_reports_hit_rates() -> None:
    expected = SearchResult(chunk_id="chunk_expected", source_file_id="file_expected", text="hit", score=1.0)
    miss = SearchResult(chunk_id="chunk_miss", source_file_id="file_miss", text="miss", score=0.1)

    metrics = RagRetriever.evaluate_retrieval(
        [
            {"id": "case_1", "expectedChunkId": "chunk_expected"},
            {"id": "case_2", "expectedSourceFileId": "file_expected"},
        ],
        {
            "case_1": [expected],
            "case_2": [miss, expected],
        },
    )

    assert metrics == {"total": 2, "top1HitRate": 0.5, "topKHitRate": 1.0}


def test_provider_info_reports_local_and_external_placeholders() -> None:
    info = build_provider_info(
        active_provider="local",
        qdrant_url="http://qdrant.local",
        pgvector_dsn=None,
        milvus_uri="http://milvus.local",
    )

    providers = {provider["name"]: provider for provider in info["providers"]}
    assert info["activeProvider"] == "local"
    assert providers["local"]["available"] is True
    assert providers["local"]["active"] is True
    assert providers["qdrant"]["configured"] is True
    assert providers["pgvector"]["configured"] is False
    assert providers["milvus"]["mode"] == "placeholder"


@pytest.mark.asyncio
async def test_rag_concurrency_guard_rejects_when_full() -> None:
    guard = RagConcurrencyGuard(limit=1)

    async with guard.slot():
        with pytest.raises(RuntimeError, match="concurrency limit"):
            async with guard.slot():
                pass

    assert guard.active == 0


def test_rag_admin_stats_and_cleanup_endpoints(monkeypatch: pytest.MonkeyPatch) -> None:
    from app import auth
    from app.main import app

    service = RagIngestionService()

    async def ingest() -> str:
        response = await service.ingest_files([("expired.md", b"Expired temporary content.")])
        return response.file_set_id

    file_set_id = asyncio.run(ingest())
    service._records[file_set_id].expires_at = datetime.now(UTC) - timedelta(hours=1)
    monkeypatch.setattr(auth, "get_api_key", lambda: None)
    monkeypatch.setattr(rag_router, "rag_ingestion_service", service)

    client = TestClient(app)
    provider_response = client.get("/rag/admin/provider-info")
    stats_response = client.get("/rag/admin/stats")
    cleanup_response = client.post("/rag/admin/cleanup")
    stats_after_cleanup = client.get("/rag/admin/stats")

    assert provider_response.status_code == 200
    assert {provider["name"] for provider in provider_response.json()["providers"]} == {
        "local",
        "qdrant",
        "pgvector",
        "milvus",
    }
    assert stats_response.status_code == 200
    assert stats_response.json()["fileSets"] == 1
    assert stats_response.json()["indexedChunks"] >= 1
    assert stats_response.json()["provider"] == "local"
    assert cleanup_response.status_code == 200
    assert cleanup_response.json() == {"cleanedFileSets": 1}
    assert stats_after_cleanup.json()["fileSets"] == 0


@pytest.mark.asyncio
async def test_rag_stream_missing_sources_returns_error() -> None:
    request = RagStreamRequest(message="没有来源时应该报错")

    stream = [event async for event in rag_router._generate_rag_stream(request)]
    rendered = "".join(stream)

    assert "event: error" in rendered
    assert "missing_sources" in rendered
    assert "event: retrieval_start" not in rendered


@pytest.mark.asyncio
async def test_rag_stream_sse_with_fake_agent(monkeypatch: pytest.MonkeyPatch) -> None:
    service = RagIngestionService()
    response = await service.ingest_files(
        [("policy.md", b"# Refund\n\nRefunds are available within 30 days.")],
        conversation_id="conv_stream",
    )
    retriever = RagRetriever(vector_store=service.get_vector_store(), embedder=service.embedder)
    monkeypatch.setattr(rag_router, "rag_tool_service", RagToolService(retriever=retriever))
    monkeypatch.setattr(
        rag_router,
        "rag_agent_runner",
        RagAgentRunner(config=RagAgentRunnerConfig()),
    )
    monkeypatch.setattr(rag_router, "create_rag_mcp_server", lambda *args, **kwargs: object())
    captured: dict[str, object] = {}

    async def fake_query_stream(**kwargs):  # noqa: ANN202
        captured.update(kwargs)
        yield AgentEvent(
            type="content_block_delta",
            subtype="text_delta",
            data={"text": "我会通过 RAG 工具查询后回答，退款期限是 30 天。"},
            conversation_id=kwargs.get("conversation_id"),
        )

    monkeypatch.setattr(rag_router.agent_service, "query_stream", fake_query_stream)

    request = RagStreamRequest(
        message="退款期限是多久？",
        conversationId="conv_stream",
        fileSetId=response.file_set_id,
    )
    stream = [event async for event in rag_router._generate_rag_stream(request)]
    rendered = "".join(stream)

    assert captured["prompt"] == "退款期限是多久？"
    assert captured["allowed_tools"] == rag_router.RAG_MCP_ALLOWED_TOOLS
    assert list(captured["mcp_servers"].keys()) == ["rag"]
    assert "event: retrieval_start" not in rendered
    assert "event: retrieval_result" not in rendered
    assert "event: agent_delta" in rendered
    assert "event: result" in rendered
    assert "30 天" in rendered


@pytest.mark.asyncio
async def test_rag_stream_does_not_pre_retrieve_before_agent(
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    service = RagIngestionService()
    response = await service.ingest_files(
        [("policy.md", b"# Refund\n\nRefunds are available within 30 days.")],
    )
    retriever = RagRetriever(vector_store=service.get_vector_store(), embedder=service.embedder)
    tool_service = RagToolService(retriever=retriever)
    monkeypatch.setattr(rag_router, "rag_tool_service", tool_service)
    monkeypatch.setattr(
        rag_router,
        "rag_agent_runner",
        RagAgentRunner(config=RagAgentRunnerConfig()),
    )
    monkeypatch.setattr(rag_router, "create_rag_mcp_server", lambda *args, **kwargs: object())

    async def unexpected_hybrid_search(**kwargs):  # noqa: ANN202, ARG001
        raise AssertionError("router must not pre-retrieve before invoking agent")

    monkeypatch.setattr(tool_service, "hybrid_search", unexpected_hybrid_search)

    async def fake_query_stream(**kwargs):  # noqa: ANN202, ARG001
        yield AgentEvent(
            type="content_block_delta",
            subtype="text_delta",
            data={"text": "通过 RAG MCP 工具回答。"},
        )

    monkeypatch.setattr(rag_router.agent_service, "query_stream", fake_query_stream)

    request = RagStreamRequest(message="退款期限是多久？", fileSetId=response.file_set_id)
    stream = [event async for event in rag_router._generate_rag_stream(request)]
    rendered = "".join(stream)

    assert "event: retrieval_start" not in rendered
    assert "event: retrieval_result" not in rendered
    assert "event: result" in rendered


@pytest.mark.asyncio
async def test_rag_stream_agent_error_stops_before_result(monkeypatch: pytest.MonkeyPatch) -> None:
    service = RagIngestionService()
    response = await service.ingest_files(
        [("policy.md", b"# Refund\n\nRefunds are available within 30 days.")],
    )
    retriever = RagRetriever(vector_store=service.get_vector_store(), embedder=service.embedder)
    monkeypatch.setattr(rag_router, "rag_tool_service", RagToolService(retriever=retriever))
    monkeypatch.setattr(
        rag_router,
        "rag_agent_runner",
        RagAgentRunner(config=RagAgentRunnerConfig()),
    )
    monkeypatch.setattr(rag_router, "create_rag_mcp_server", lambda *args, **kwargs: object())

    async def fake_query_stream(**kwargs):  # noqa: ANN202, ARG001
        yield AgentEvent(type="error", data="agent failed")

    monkeypatch.setattr(rag_router.agent_service, "query_stream", fake_query_stream)

    request = RagStreamRequest(message="退款期限是多久？", fileSetId=response.file_set_id)
    stream = [event async for event in rag_router._generate_rag_stream(request)]
    rendered = "".join(stream)

    assert "event: error" in rendered
    assert "agent_error" in rendered
    assert "event: result" not in rendered
